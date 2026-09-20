import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, Any
from backend.models import FinancialMetrics, DealInputParams
from backend.merger_engine import calculate_merger_math

def build_excel_deal_model(
    acquirer: FinancialMetrics,
    target: FinancialMetrics,
    params: DealInputParams,
    deal_results: Dict[str, Any]
) -> io.BytesIO:
    """
    Generates an institutional Wall Street M&A model Excel file (.xlsx) with active Excel formulas.
    Formatting rules:
    - Input hardcoded cells: Blue text (#002060), soft fill
    - Dynamic formula cells: Black text (#000000), formatted numbers
    - Headers: Dark Navy fill (#1F4E79) with white bold text
    - Includes multi-tab coverage: Summary & Model, 2D Sensitivity Matrix
    """
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # STYLES
    # -------------------------------------------------------------
    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=11, bold=True, color="1F4E79")
    
    font_input = Font(name="Calibri", size=11, bold=True, color="002060") # IB Blue text for inputs
    font_formula = Font(name="Calibri", size=11, bold=False, color="000000") # Black for formulas
    font_bold_formula = Font(name="Calibri", size=11, bold=True, color="000000")
    
    fill_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_input = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    fill_accent = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_accretive = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Soft green
    fill_dilutive = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid") # Soft red

    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")

    thin_border_side = Side(border_style="thin", color="D9D9D9")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_total = Border(top=thin_border_side, bottom=Side(border_style="double", color="000000"))

    # Number Formats
    FMT_CURRENCY = "$#,##0"
    FMT_PRICE = "$#,##0.00"
    FMT_SHARES = "#,##0"
    FMT_PCT = "0.0%"
    FMT_PCT_EXACT = "0.00%"
    FMT_EPS = "$#,##0.00"

    # -------------------------------------------------------------
    # TAB 1: DEAL SUMMARY & CONSOLIDATION MODEL
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Deal Model"
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:D1")
    ws1["A1"] = f"SINERGIA M&A ACCRETION / DILUTION MODEL: {acquirer.ticker} / {target.ticker}"
    ws1["A1"].font = font_title
    ws1["A1"].fill = fill_header
    ws1["A1"].alignment = align_center

    # Section 1: Standalone Inputs
    ws1.cell(row=3, column=1, value="1. STANDALONE COMPANY FINANCIALS").font = font_section
    ws1.cell(row=3, column=3, value=acquirer.ticker).font = font_header
    ws1.cell(row=3, column=3).fill = fill_header
    ws1.cell(row=3, column=3).alignment = align_center
    ws1.cell(row=3, column=4, value=target.ticker).font = font_header
    ws1.cell(row=3, column=4).fill = fill_header
    ws1.cell(row=3, column=4).alignment = align_center

    row_data = [
        ("Current Share Price", acquirer.share_price, target.share_price, FMT_PRICE),
        ("Diluted Shares Outstanding", acquirer.diluted_shares, target.diluted_shares, FMT_SHARES),
        ("Net Income ($)", acquirer.net_income, target.net_income, FMT_CURRENCY),
        ("Book Value of Net Assets ($)", acquirer.book_value_net_assets, target.book_value_net_assets, FMT_CURRENCY),
        ("Effective Tax Rate", acquirer.effective_tax_rate, target.effective_tax_rate, FMT_PCT_EXACT),
    ]

    r = 4
    acq_price_cell = f"C4"
    acq_shares_cell = f"C5"
    acq_ni_cell = f"C6"
    tgt_price_cell = f"D4"
    tgt_shares_cell = f"D5"
    tgt_ni_cell = f"D6"
    tgt_bv_cell = f"D7"
    tax_rate_cell = f"C8"

    for label, v1, v2, fmt in row_data:
        ws1.cell(row=r, column=1, value=label).font = font_formula
        
        c1 = ws1.cell(row=r, column=3, value=v1)
        c1.font = font_input
        c1.fill = fill_input
        c1.number_format = fmt
        c1.alignment = align_right
        
        c2 = ws1.cell(row=r, column=4, value=v2)
        c2.font = font_input
        c2.fill = fill_input
        c2.number_format = fmt
        c2.alignment = align_right
        r += 1

    # Section 2: Transaction Structuring Inputs
    r += 1
    ws1.cell(row=r, column=1, value="2. TRANSACTION STRUCTURING ASSUMPTIONS").font = font_section
    r += 1

    ws1.cell(row=r, column=1, value="Offer Premium %").font = font_formula
    c = ws1.cell(row=r, column=3, value=params.offer_premium_pct / 100.0)
    c.font = font_input; c.fill = fill_input; c.number_format = FMT_PCT; c.alignment = align_right
    prem_cell = f"C{r}"
    r += 1

    ws1.cell(row=r, column=1, value="% Cash Consideration").font = font_formula
    c = ws1.cell(row=r, column=3, value=params.cash_pct / 100.0)
    c.font = font_input; c.fill = fill_input; c.number_format = FMT_PCT; c.alignment = align_right
    cash_pct_cell = f"C{r}"
    r += 1

    ws1.cell(row=r, column=1, value="% Debt Consideration").font = font_formula
    c = ws1.cell(row=r, column=3, value=params.debt_pct / 100.0)
    c.font = font_input; c.fill = fill_input; c.number_format = FMT_PCT; c.alignment = align_right
    debt_pct_cell = f"C{r}"
    r += 1

    ws1.cell(row=r, column=1, value="% Stock Consideration").font = font_formula
    c = ws1.cell(row=r, column=3, value=params.stock_pct / 100.0)
    c.font = font_input; c.fill = fill_input; c.number_format = FMT_PCT; c.alignment = align_right
    stock_pct_cell = f"C{r}"
    r += 1

    ws1.cell(row=r, column=1, value="Cost of Debt %").font = font_formula
    c = ws1.cell(row=r, column=3, value=params.cost_of_debt_pct / 100.0)
    c.font = font_input; c.fill = fill_input; c.number_format = FMT_PCT_EXACT; c.alignment = align_right
    cost_debt_cell = f"C{r}"
    r += 1

    ws1.cell(row=r, column=1, value="Foregone Interest Rate on Cash %").font = font_formula
    c = ws1.cell(row=r, column=3, value=params.foregone_interest_rate_pct / 100.0)
    c.font = font_input; c.fill = fill_input; c.number_format = FMT_PCT_EXACT; c.alignment = align_right
    interest_cash_cell = f"C{r}"
    r += 1

    ws1.cell(row=r, column=1, value="Annual Pre-Tax Synergies ($)").font = font_formula
    c = ws1.cell(row=r, column=3, value=params.synergies_pre_tax)
    c.font = font_input; c.fill = fill_input; c.number_format = FMT_CURRENCY; c.alignment = align_right
    synergies_cell = f"C{r}"
    r += 1

    # Section 3: Dynamic Transaction Mechanics & Valuation (ACTIVE EXCEL FORMULAS)
    r += 1
    ws1.cell(row=r, column=1, value="3. VALUATION & BALANCE SHEET ADJUSTMENTS").font = font_section
    r += 1

    # Offer Price = Target Price * (1 + Premium %)
    ws1.cell(row=r, column=1, value="Offer Price Per Target Share").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"={tgt_price_cell}*(1+{prem_cell})")
    c.font = font_formula; c.number_format = FMT_PRICE; c.alignment = align_right
    offer_price_cell = f"C{r}"
    r += 1

    # Purchase Equity Value = Offer Price * Target Shares
    ws1.cell(row=r, column=1, value="Purchase Equity Value").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"={offer_price_cell}*{tgt_shares_cell}")
    c.font = font_bold_formula; c.number_format = FMT_CURRENCY; c.alignment = align_right
    eq_val_cell = f"C{r}"
    r += 1

    # Goodwill Created = MAX(0, Equity Value - Book Value Net Assets)
    ws1.cell(row=r, column=1, value="Goodwill Created").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"=MAX(0, {eq_val_cell}-{tgt_bv_cell})")
    c.font = font_formula; c.number_format = FMT_CURRENCY; c.alignment = align_right
    goodwill_cell = f"C{r}"
    r += 1

    # Cash Used = Equity Value * Cash %
    ws1.cell(row=r, column=1, value="Cash Consideration Used").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"={eq_val_cell}*{cash_pct_cell}")
    c.font = font_formula; c.number_format = FMT_CURRENCY; c.alignment = align_right
    cash_used_cell = f"C{r}"
    r += 1

    # Debt Raised = Equity Value * Debt %
    ws1.cell(row=r, column=1, value="Debt Raised").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"={eq_val_cell}*{debt_pct_cell}")
    c.font = font_formula; c.number_format = FMT_CURRENCY; c.alignment = align_right
    debt_raised_cell = f"C{r}"
    r += 1

    # Stock Issued Value = Equity Value * Stock %
    ws1.cell(row=r, column=1, value="Stock Consideration Value").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"={eq_val_cell}*{stock_pct_cell}")
    c.font = font_formula; c.number_format = FMT_CURRENCY; c.alignment = align_right
    stock_val_cell = f"C{r}"
    r += 1

    # New Shares Issued = Stock Value / Acquirer Price
    ws1.cell(row=r, column=1, value="New Shares Issued").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"={stock_val_cell}/{acq_price_cell}")
    c.font = font_formula; c.number_format = FMT_SHARES; c.alignment = align_right
    new_shares_cell = f"C{r}"
    r += 1

    # Pro-Forma Shares = Acquirer Shares + New Shares
    ws1.cell(row=r, column=1, value="Pro-Forma Shares Outstanding").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"={acq_shares_cell}+{new_shares_cell}")
    c.font = font_bold_formula; c.number_format = FMT_SHARES; c.alignment = align_right
    pf_shares_cell = f"C{r}"
    r += 1

    # Section 4: Income Consolidation & EPS Accretion/Dilution
    r += 1
    ws1.cell(row=r, column=1, value="4. INCOME STATEMENT CONSOLIDATION & EPS IMPACT").font = font_section
    r += 1

    # Net Interest Burden = (Debt Raised * Cost of Debt) + (Cash Used * Interest Rate)
    ws1.cell(row=r, column=1, value="Net Pre-Tax Interest Burden").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"=({debt_raised_cell}*{cost_debt_cell})+({cash_used_cell}*{interest_cash_cell})")
    c.font = font_formula; c.number_format = FMT_CURRENCY; c.alignment = align_right
    interest_burden_cell = f"C{r}"
    r += 1

    # Post-Tax Synergy Adj = (Pre-Tax Synergies - Interest Burden) * (1 - Tax Rate)
    ws1.cell(row=r, column=1, value="Post-Tax Net Synergy Adjustment").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"=({synergies_cell}-{interest_burden_cell})*(1-{tax_rate_cell})")
    c.font = font_formula; c.number_format = FMT_CURRENCY; c.alignment = align_right
    post_tax_syn_cell = f"C{r}"
    r += 1

    # Pro-Forma Net Income = Acquirer NI + Target NI + Post-Tax Synergies
    ws1.cell(row=r, column=1, value="Pro-Forma Combined Net Income").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"={acq_ni_cell}+{tgt_ni_cell}+{post_tax_syn_cell}")
    c.font = font_bold_formula; c.number_format = FMT_CURRENCY; c.alignment = align_right
    pf_ni_cell = f"C{r}"
    r += 1

    # Standalone EPS = Acquirer NI / Acquirer Shares
    ws1.cell(row=r, column=1, value="Acquirer Standalone EPS").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"={acq_ni_cell}/{acq_shares_cell}")
    c.font = font_bold_formula; c.number_format = FMT_EPS; c.alignment = align_right
    standalone_eps_cell = f"C{r}"
    r += 1

    # Pro-Forma EPS = Pro-Forma Net Income / Pro-Forma Shares
    ws1.cell(row=r, column=1, value="Pro-Forma EPS").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"={pf_ni_cell}/{pf_shares_cell}")
    c.font = font_bold_formula; c.number_format = FMT_EPS; c.alignment = align_right
    pf_eps_cell = f"C{r}"
    r += 1

    # Accretion / Dilution % = (Pro-Forma EPS - Standalone EPS) / ABS(Standalone EPS)
    ws1.cell(row=r, column=1, value="EPS Accretion / (Dilution) %").font = font_section
    c = ws1.cell(row=r, column=3, value=f"=({pf_eps_cell}-{standalone_eps_cell})/ABS({standalone_eps_cell})")
    c.font = Font(name="Calibri", size=12, bold=True, color="002060")
    c.fill = fill_accretive if deal_results["accretion_dilution_pct"] >= 0 else fill_dilutive
    c.number_format = FMT_PCT_EXACT; c.alignment = align_right; c.border = border_total
    acc_pct_excel_cell = f"C{r}"
    r += 1

    # Breakeven Synergies Solver Formula = ((Standalone EPS * Pro-Forma Shares) - Acq NI - Tgt NI)/(1 - Tax Rate) + Interest Burden
    ws1.cell(row=r, column=1, value="Required Breakeven Pre-Tax Synergies").font = font_formula
    c = ws1.cell(row=r, column=3, value=f"=((({standalone_eps_cell}*{pf_shares_cell})-{acq_ni_cell}-{tgt_ni_cell})/(1-{tax_rate_cell}))+{interest_burden_cell}")
    c.font = font_bold_formula; c.number_format = FMT_CURRENCY; c.alignment = align_right; c.fill = fill_accent
    r += 1

    # Adjust Column Widths
    ws1.column_dimensions["A"].width = 42
    ws1.column_dimensions["B"].width = 5
    ws1.column_dimensions["C"].width = 25
    ws1.column_dimensions["D"].width = 25

    # -------------------------------------------------------------
    # TAB 2: 2D SENSITIVITY MATRIX
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Sensitivity Matrix")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:G1")
    ws2["A1"] = "2D ACCRETION / (DILUTION) % SENSITIVITY MATRIX"
    ws2["A1"].font = font_title; ws2["A1"].fill = fill_header; ws2["A1"].alignment = align_center

    ws2.cell(row=3, column=1, value="Y-Axis: Offer Premium % | X-Axis: % Stock Consideration").font = font_section

    # Headers for Stock Consideration (0%, 25%, 50%, 75%, 100%)
    stock_range = [0.0, 0.25, 0.50, 0.75, 1.00]
    premium_range = [0.10, 0.20, 0.30, 0.40, 0.50]

    ws2.cell(row=5, column=2, value="Offer Premium \\ Stock %").font = font_header
    ws2.cell(row=5, column=2).fill = fill_header; ws2.cell(row=5, column=2).alignment = align_center

    col_idx = 3
    for s in stock_range:
        c = ws2.cell(row=5, column=col_idx, value=s)
        c.font = font_header; c.fill = fill_header; c.number_format = FMT_PCT; c.alignment = align_center
        col_idx += 1

    r_idx = 6
    for prem in premium_range:
        c_prem = ws2.cell(row=r_idx, column=2, value=prem)
        c_prem.font = font_header; c_prem.fill = fill_header; c_prem.number_format = FMT_PCT; c_prem.alignment = align_center
        
        c_col = 3
        for stock in stock_range:
            # Create standard parameter copy result
            p_copy = params.model_copy()
            p_copy.offer_premium_pct = prem * 100.0
            p_copy.stock_pct = stock * 100.0
            rem = 100.0 - (stock * 100.0)
            p_copy.cash_pct = rem
            p_copy.debt_pct = 0.0
            
            sub_res = calculate_merger_math(acquirer, target, p_copy)
            acc_val = sub_res["accretion_dilution_pct"] / 100.0
            
            c_cell = ws2.cell(row=r_idx, column=c_col, value=acc_val)
            c_cell.font = font_formula
            c_cell.number_format = FMT_PCT_EXACT
            c_cell.alignment = align_right
            c_cell.border = border_cell
            c_cell.fill = fill_accretive if acc_val >= 0 else fill_dilutive
            
            c_col += 1
        r_idx += 1

    ws2.column_dimensions["A"].width = 5
    ws2.column_dimensions["B"].width = 25
    for c_i in range(3, 8):
        ws2.column_dimensions[get_column_letter(c_i)].width = 18

    # Save to BytesIO stream
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
